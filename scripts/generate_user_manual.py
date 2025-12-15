"""
Markdown 매뉴얼을 Word(.docx) 및 PDF로 변환하는 유틸리티.

목표:
- 제출용 매뉴얼을 자동으로 생성(WORD/PDF)
- 한글 PDF 깨짐 방지를 위해 Windows 기본 한글 폰트(맑은 고딕)를 우선 사용

사용법(프로젝트 루트에서):
  1) (최초 1회) pip install -r requirements-manual.txt
  2) python scripts/generate_user_manual.py

출력:
  - docs/BIMO_사용자매뉴얼.docx
  - docs/BIMO_사용자매뉴얼.pdf
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import List, Optional, Tuple


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_MD = ROOT / "docs" / "BIMO_사용자매뉴얼.md"
DEFAULT_DOCX = ROOT / "docs" / "BIMO_사용자매뉴얼.docx"
DEFAULT_PDF = ROOT / "docs" / "BIMO_사용자매뉴얼.pdf"


@dataclass
class Block:
    kind: str  # heading|paragraph|bullets|code|table|hr
    level: int = 0
    text: str = ""
    bullets: Optional[List[str]] = None
    code: Optional[str] = None
    table: Optional[List[List[str]]] = None  # rows


def _is_table_separator(line: str) -> bool:
    s = line.strip()
    if "|" not in s:
        return False
    # --- | --- 같은 형태
    parts = [p.strip() for p in s.strip("|").split("|")]
    if not parts:
        return False
    return all(p and set(p) <= {"-", ":", " "} for p in parts)


def parse_markdown(md_text: str) -> List[Block]:
    """
    매우 단순한 Markdown 파서:
    - #/##/### heading
    - - bullet
    - ``` code fence
    - | ... | table (header + separator + rows)
    - --- horizontal rule
    - 그 외는 paragraph
    """
    lines = md_text.splitlines()
    blocks: List[Block] = []

    i = 0
    paragraph_buf: List[str] = []

    def flush_paragraph() -> None:
        nonlocal paragraph_buf
        text = "\n".join([l.rstrip() for l in paragraph_buf]).strip()
        if text:
            blocks.append(Block(kind="paragraph", text=text))
        paragraph_buf = []

    while i < len(lines):
        line = lines[i].rstrip("\n")

        # code fence
        if line.strip().startswith("```"):
            flush_paragraph()
            fence = line.strip()
            i += 1
            code_lines: List[str] = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i].rstrip("\n"))
                i += 1
            # closing fence consume
            if i < len(lines) and lines[i].strip().startswith("```"):
                i += 1
            blocks.append(Block(kind="code", code="\n".join(code_lines)))
            continue

        # horizontal rule
        if line.strip() == "---":
            flush_paragraph()
            blocks.append(Block(kind="hr"))
            i += 1
            continue

        # heading
        if line.startswith("#"):
            flush_paragraph()
            hashes = len(line) - len(line.lstrip("#"))
            title = line[hashes:].strip()
            blocks.append(Block(kind="heading", level=hashes, text=title))
            i += 1
            continue

        # table (markdown)
        if "|" in line and i + 1 < len(lines) and _is_table_separator(lines[i + 1]):
            flush_paragraph()
            header = [c.strip() for c in line.strip().strip("|").split("|")]
            i += 2  # skip header + separator
            rows: List[List[str]] = [header]
            while i < len(lines):
                row_line = lines[i].strip()
                if not row_line or "|" not in row_line:
                    break
                row = [c.strip() for c in row_line.strip().strip("|").split("|")]
                # 열 개수 맞추기
                while len(row) < len(header):
                    row.append("")
                rows.append(row[: len(header)])
                i += 1
            blocks.append(Block(kind="table", table=rows))
            continue

        # bullets
        if line.strip().startswith("- "):
            flush_paragraph()
            items: List[str] = []
            while i < len(lines) and lines[i].strip().startswith("- "):
                items.append(lines[i].strip()[2:].strip())
                i += 1
            blocks.append(Block(kind="bullets", bullets=items))
            continue

        # blank line: paragraph boundary
        if not line.strip():
            flush_paragraph()
            i += 1
            continue

        paragraph_buf.append(line)
        i += 1

    flush_paragraph()
    return blocks


def render_docx(blocks: List[Block], out_path: Path) -> None:
    from docx import Document
    from docx.shared import Pt

    doc = Document()

    # 기본 폰트 크기(Word 기본이 충분하지만, 코드/가독성을 위해 조금 정리)
    normal_style = doc.styles["Normal"]
    normal_style.font.name = "맑은 고딕"
    normal_style.font.size = Pt(11)

    for b in blocks:
        if b.kind == "heading":
            level = max(1, min(4, b.level))  # docx는 1~9 지원, 여기선 1~4만 사용
            doc.add_heading(b.text, level=level)
        elif b.kind == "paragraph":
            for para in b.text.split("\n"):
                doc.add_paragraph(para.strip())
        elif b.kind == "bullets" and b.bullets:
            for item in b.bullets:
                doc.add_paragraph(item, style="List Bullet")
        elif b.kind == "code" and b.code is not None:
            p = doc.add_paragraph()
            run = p.add_run(b.code)
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        elif b.kind == "table" and b.table:
            rows = b.table
            table = doc.add_table(rows=0, cols=len(rows[0]))
            table.style = "Table Grid"
            for r_idx, row in enumerate(rows):
                tr = table.add_row().cells
                for c_idx, cell_text in enumerate(row):
                    tr[c_idx].text = cell_text
            # 헤더 굵게 등은 생략(필요 시 수동 편집)
        elif b.kind == "hr":
            doc.add_paragraph("―" * 20)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path.as_posix())


def _find_windows_korean_font() -> Optional[Tuple[str, str]]:
    """
    reportlab PDF 한글 깨짐 방지를 위해 TTF 경로를 찾습니다.
    Returns: (font_name, ttf_path)
    """
    candidates = [
        ("MalgunGothic", r"C:\Windows\Fonts\malgun.ttf"),
        ("MalgunGothicBold", r"C:\Windows\Fonts\malgunbd.ttf"),
        ("Batang", r"C:\Windows\Fonts\batang.ttf"),
        ("Gulim", r"C:\Windows\Fonts\gulim.ttf"),
    ]
    for name, path in candidates:
        if Path(path).exists():
            return name, path
    return None


def render_pdf(blocks: List[Block], out_path: Path) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import (
        SimpleDocTemplate,
        Paragraph,
        Spacer,
        Preformatted,
        Table,
        TableStyle,
    )
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    styles = getSampleStyleSheet()

    font_info = _find_windows_korean_font()
    base_font = "Helvetica"
    mono_font = "Courier"
    if font_info:
        font_name, font_path = font_info
        try:
            pdfmetrics.registerFont(TTFont(font_name, font_path))
            base_font = font_name
            mono_font = font_name  # 코드도 한글 포함 가능하므로 동일 폰트 사용
        except Exception:
            # 폰트 등록 실패 시 기본 폰트로 진행(한글은 깨질 수 있음)
            pass

    h1 = ParagraphStyle("H1", parent=styles["Heading1"], fontName=base_font)
    h2 = ParagraphStyle("H2", parent=styles["Heading2"], fontName=base_font)
    h3 = ParagraphStyle("H3", parent=styles["Heading3"], fontName=base_font)
    body = ParagraphStyle("Body", parent=styles["BodyText"], fontName=base_font, leading=16)
    bullet = ParagraphStyle("Bullet", parent=styles["BodyText"], fontName=base_font, leftIndent=14, leading=16)
    code_style = ParagraphStyle("Code", parent=styles["Code"], fontName=mono_font, fontSize=9, leading=12)

    doc = SimpleDocTemplate(
        out_path.as_posix(),
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title="BIMO-BE 사용자 매뉴얼",
    )

    story = []
    for b in blocks:
        if b.kind == "heading":
            if b.level <= 1:
                story.append(Paragraph(b.text, h1))
            elif b.level == 2:
                story.append(Paragraph(b.text, h2))
            else:
                story.append(Paragraph(b.text, h3))
            story.append(Spacer(1, 6))
        elif b.kind == "paragraph":
            for para in b.text.split("\n"):
                if para.strip():
                    story.append(Paragraph(para.strip(), body))
            story.append(Spacer(1, 6))
        elif b.kind == "bullets" and b.bullets:
            for item in b.bullets:
                story.append(Paragraph(f"• {item}", bullet))
            story.append(Spacer(1, 6))
        elif b.kind == "code" and b.code is not None:
            story.append(Preformatted(b.code, code_style))
            story.append(Spacer(1, 6))
        elif b.kind == "table" and b.table:
            data = b.table
            t = Table(data, hAlign="LEFT")
            t.setStyle(
                TableStyle(
                    [
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.whitesmoke),
                        ("FONTNAME", (0, 0), (-1, -1), base_font),
                        ("FONTSIZE", (0, 0), (-1, -1), 9.5),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 4),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                        ("TOPPADDING", (0, 0), (-1, -1), 3),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                    ]
                )
            )
            story.append(t)
            story.append(Spacer(1, 8))
        elif b.kind == "hr":
            story.append(Paragraph("―" * 30, body))
            story.append(Spacer(1, 8))

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.build(story)


def main() -> None:
    md_path = DEFAULT_MD
    out_docx = DEFAULT_DOCX
    out_pdf = DEFAULT_PDF

    if not md_path.exists():
        raise FileNotFoundError(f"Markdown 파일을 찾을 수 없습니다: {md_path}")

    md_text = md_path.read_text(encoding="utf-8")
    blocks = parse_markdown(md_text)

    render_docx(blocks, out_docx)
    render_pdf(blocks, out_pdf)

    print("[OK] 매뉴얼 생성 완료")
    print(f"- DOCX: {out_docx}")
    print(f"- PDF : {out_pdf}")


if __name__ == "__main__":
    main()



